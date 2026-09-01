from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
PLAIN_TEXT = Path(sys.argv[3])


def set_font(run, name: str, size: float, *, bold: bool = False, italic: bool = False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_hyperlink(paragraph, label: str, url: str):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run_element = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_fonts = OxmlElement("w:rFonts")
    run_fonts.set(qn("w:ascii"), "Arial")
    run_fonts.set(qn("w:hAnsi"), "Arial")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "22")
    run_properties.extend([run_fonts, color, underline, size])
    text = OxmlElement("w:t")
    text.text = label
    run_element.extend([run_properties, text])
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(\[[^\]]+\]\(https?://[^)]+\)|\*\*[^*]+\*\*|\*[^*]+\*)"
)


def add_inline(paragraph, text: str, *, base_italic: bool = False):
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_font(run, "Arial", 11, italic=base_italic)
        token = match.group(0)
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", token)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, "Arial", 11, bold=True, italic=base_italic)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, "Arial", 11, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_font(run, "Arial", 11, italic=base_italic)


raw = SOURCE.read_text(encoding="utf-8")
subtitle_match = re.search(r"## Subtitle\s*\n\s*(.+)", raw)
subtitle = subtitle_match.group(1).strip() if subtitle_match else ""

parts = raw.split("\n---\n")
if len(parts) < 2:
    raise SystemExit("Expected publishing metadata followed by article body.")
body = "\n---\n".join(parts[1:]).strip()
lines = body.splitlines()

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.right_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, before, after, color in [
    ("Heading 1", 20, 20, 6, "000000"),
    ("Heading 2", 16, 18, 6, "000000"),
    ("Heading 3", 14, 16, 4, "434343"),
]:
    style = doc.styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = False
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

in_code = False
code_lines: list[str] = []
title_added = False
plain_lines: list[str] = []

for raw_line in lines:
    line = raw_line.rstrip()
    if line.startswith("```"):
        if in_code:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.22)
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.line_spacing = 1.0
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F3F4F6")
            paragraph._p.get_or_add_pPr().append(shading)
            run = paragraph.add_run("\n".join(code_lines))
            set_font(run, "Courier New", 9)
            plain_lines.extend(code_lines)
            code_lines = []
            in_code = False
        else:
            in_code = True
        continue
    if in_code:
        code_lines.append(line)
        continue
    if not line:
        plain_lines.append("")
        continue
    if line == "---":
        continue
    if line.startswith("# ") and not title_added:
        title_text = line[2:].strip()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(title_text)
        set_font(run, "Arial", 26)
        title_added = True
        plain_lines.append(title_text)
        if subtitle:
            subtitle_paragraph = doc.add_paragraph()
            subtitle_paragraph.paragraph_format.space_before = Pt(0)
            subtitle_paragraph.paragraph_format.space_after = Pt(18)
            subtitle_run = subtitle_paragraph.add_run(subtitle)
            set_font(subtitle_run, "Arial", 13)
            subtitle_run.font.color.rgb = RGBColor(85, 85, 85)
            plain_lines.extend([subtitle, ""])
        continue
    if line.startswith("## "):
        text = line[3:].strip()
        doc.add_paragraph(text, style="Heading 2")
        plain_lines.extend(["", text, ""])
        continue
    if line.startswith("### "):
        text = line[4:].strip()
        doc.add_paragraph(text, style="Heading 3")
        plain_lines.extend(["", text, ""])
        continue
    if line.startswith("> "):
        text = line[2:].strip()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.right_indent = Inches(0.25)
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.paragraph_format.line_spacing = 1.15
        border = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "10")
        left.set(qn("w:color"), "DADCE0")
        border.append(left)
        paragraph._p.get_or_add_pPr().append(border)
        add_inline(paragraph, text, base_italic=True)
        plain_lines.append(text)
        continue
    if line.startswith("- "):
        text = line[2:].strip()
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        add_inline(paragraph, text)
        plain_lines.append(f"• {re.sub(INLINE_PATTERN, lambda m: m.group(0).strip('*'), text)}")
        continue
    numbered = re.match(r"^(\d+)\.\s+(.+)", line)
    if numbered:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        add_inline(paragraph, numbered.group(2))
        plain_lines.append(f"{numbered.group(1)}. {numbered.group(2)}")
        continue

    italic_disclosure = line.startswith("*Disclosure:") and line.endswith("*")
    paragraph = doc.add_paragraph()
    if italic_disclosure:
        line = line[1:-1]
        paragraph.paragraph_format.space_before = Pt(16)
    add_inline(paragraph, line, base_italic=italic_disclosure)
    plain_line = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", line)
    plain_line = plain_line.replace("**", "").replace("*", "")
    plain_lines.append(plain_line)

doc.core_properties.title = "The Hardest Part of Building an AI News Reader Wasn’t the AI"
doc.core_properties.subject = "Substack-ready article"
doc.core_properties.author = "Tarun Tomar"
doc.core_properties.keywords = "AI news reader, software engineering, news aggregation, Trace"
doc.save(OUTPUT)

while plain_lines and not plain_lines[-1].strip():
    plain_lines.pop()
PLAIN_TEXT.write_text("\n".join(plain_lines) + "\n", encoding="utf-8")
